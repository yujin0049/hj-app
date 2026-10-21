"""T053 (261021) 사물함 장부 xlsx → data/lockers.json 시험. 실행: python -m unittest tests/test_lockers_xlsx.py -v
- 결과에 이름·학번 열(키)이 없고, 칸번호·열쇠 수·상태(+위치) 말고는 아무것도 안 들어가는지
- 가짜 이름·학번을 채운 장부를 넣어도 결과에 그 값이 하나도 안 나오는지
- 들어 있는 data/lockers.json 이 엔진실습실 116칸인지"""
import json, os, re, shutil, sys, tempfile, unittest
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, 'tools'))
import openpyxl
import make_lockers_xlsx as X

ALLOWED_KEYS = {'번호', '위치', '열쇠수', '상태', '시작일', '반납일'}
PII_KEY = re.compile(r'(이름|성명|학번|학생|전화|연락처|name|비고)', re.I)
FAKE_NAMES = ['홍길동', '김가짜', '이테스트']
FAKE_IDS = ['20269999001', '20269999002', '2026999903']


def fixture(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    e = wb.create_sheet('엔진실습실')
    e.append(['엔진실습실 사물함'])
    e.append(['설명'])
    e.append([])
    e.append(['칸번호', '구역', '열', '위/아래', '문 번호', '열쇠 수', '상태', '학번', '이름', '시작일', '반납일', '분실 시 복제 반납일', '자동 확인', '비고'])
    e.append(['1층-001', 'A', 1, '위', '028', 2, '빈칸', FAKE_IDS[0], FAKE_NAMES[0], '2026-03-02', '', '', '', FAKE_NAMES[1] + ' 메모'])
    e.append(['1층-002', 'A', 1, '아래', '029', 0, '사용', FAKE_IDS[1], FAKE_NAMES[1], '', '', '', '', ''])
    e.append(['2층-001', 'B', 1, '위', '?32', None, None, '', '', '', '', '', '', ''])
    s = wb.create_sheet('2층학생')
    s.append(['2층 학생 사물함'])
    s.append(['설명'])
    s.append([])
    s.append(['번호', '열쇠 수', '상태', '학번', '이름', '기존 사용자 신고일', '시작일', '반납일', '분실 시 복제 반납일', '자동 확인', '비고'])
    s.append([7, 1, '사용', FAKE_IDS[2], FAKE_NAMES[2], '2026-10-21', '', '', '', '', ''])
    s.append([None, None, None, '', '', '', '', '', '', '', ''])
    wb.save(path)


class T변환(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.x = os.path.join(self.tmp, '장부.xlsx')
        self.out = os.path.join(self.tmp, 'lockers.json')
        fixture(self.x)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_no_name_or_student_id(self):
        self.assertEqual(X.main([self.x, '--out', self.out]), 0)
        text = open(self.out, encoding='utf-8').read()
        for v in FAKE_NAMES + FAKE_IDS:
            self.assertNotIn(v, text)
        self.assertNotRegex(text, r'\d{7,}')
        d = json.loads(text)
        for it in d['items']:
            self.assertEqual(set(it), ALLOWED_KEYS)
            for k in it:
                self.assertNotRegex(k, PII_KEY)
            self.assertEqual(it['시작일'], '')   # 시작일 열도 읽지 않음
        self.assertEqual([(i['위치'], i['번호'], i['열쇠수'], i['상태']) for i in d['items']], [
            ('엔진실습실 1층', '1', 2, '빈칸'), ('엔진실습실 1층', '2', 0, '사용 중지'),
            ('엔진실습실 2층', '1', None, '확인 필요'), ('2층 휴게실', '7', 1, '사용')])
        self.assertEqual(d['sheets']['2층학생'], {'번호 있는 칸': 1, '번호 빈 줄(넣지 않음)': 1})

    def test_only_three_columns_read(self):
        wb = openpyxl.load_workbook(self.x)
        rows, info = X.rows_from(wb)
        for r in rows:
            self.assertEqual(len(r), 4)   # 위치·번호·열쇠수·상태
            for v in r:
                self.assertNotIn(v, FAKE_NAMES + FAKE_IDS)

    def test_student_id_like_number_refused(self):
        wb = openpyxl.load_workbook(self.x)
        wb['2층학생'].cell(row=5, column=1).value = '20261234'
        wb.save(self.x)
        self.assertEqual(X.main([self.x, '--out', self.out]), 2)
        self.assertFalse(os.path.exists(self.out))


class T들어있는자료(unittest.TestCase):
    def test_committed_data(self):
        with open(os.path.join(ROOT, 'data', 'lockers.json'), encoding='utf-8') as f:
            d = json.load(f)
        self.assertEqual(len(d['items']), 116)
        self.assertEqual(sum(1 for i in d['items'] if i['위치'] == '엔진실습실 1층'), 102)
        self.assertEqual(sum(1 for i in d['items'] if i['위치'] == '엔진실습실 2층'), 14)
        for it in d['items']:
            self.assertEqual(set(it), ALLOWED_KEYS)
            self.assertIn(it['상태'], ['빈칸', '사용', '확인 필요', '사용 중지'])
        self.assertNotRegex(json.dumps(d, ensure_ascii=False), r'"(이름|학번|성명)"')


if __name__ == '__main__':
    unittest.main()
