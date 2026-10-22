"""정례 문서 달력 xlsx → data/regular_docs.json (유진님 10.22. 「사이트에도 넣으라」, 대장 전달).
실행: python3 tools/make_regular_docs.py [원본.xlsx] [출력.json]
- 「달력」 시트(분류·문서·주기·횟수 + 3월~2월 ●/○)를 읽고, 맨 아래 「그 달 건수」 합계 줄은 버려요.
- 「월별 목록」 시트와 달별 건수가 맞는지 맞춰 보고, 다르면 멈춰요(원본이 어긋난 채로 올라가지 않게).
- 공개 사이트에 올라가므로 사람 이름처럼 보이는 글자(○○○ 말고 「이름 + 선생님」 꼴)가 있으면 멈춰요."""
import json, os, re, sys
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = '/home/box/vault/봇공유/10 업무/정례문서/261022_정례문서달력_원본.xlsx'
OUT = os.path.join(ROOT, 'data', 'regular_docs.json')
MONTHS = [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2]
NAME_RE = re.compile(r'(?<![가-힣])(?!담당)[가-힣]{2,3}\s?(선생님|주무관|교수님|원장님|팀장님)')
CYCLE_ORDER = ['매월', '분기마다', '학기마다', '연 1회', '수시 반복']


def cycle_key(c):
    for i, k in enumerate(CYCLE_ORDER):
        if str(c).startswith(k):
            return i
    return len(CYCLE_ORDER)


def build(src):
    import openpyxl
    wb = openpyxl.load_workbook(src, data_only=False)
    ws = wb['달력']
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == '분류' and r[1] == '문서')
    head = rows[hi]
    mcol = {}
    for j, h in enumerate(head):
        m = re.fullmatch(r'(\d{1,2})월', str(h or ''))
        if m:
            mcol[int(m.group(1))] = j
    if sorted(mcol) != sorted(MONTHS):
        raise SystemExit('달력 시트 월 칸이 3월~2월 12개가 아니에요: %s' % sorted(mcol))
    title = str(rows[0][0] or '')
    m = re.search(r'자료 갱신 ([0-9-]+ [0-9:]+)', title)
    docs = []
    for r in rows[hi + 1:]:
        if not r or not r[0] or not r[1]:
            continue   # 빈 줄, 「그 달 건수」 합계 줄
        marks = {}
        for mo in MONTHS:
            v = str(r[mcol[mo]] or '').strip()
            if v in ('●', '○'):
                marks[str(mo)] = v
            elif v:
                raise SystemExit('알 수 없는 표시 %r: %s %d월' % (v, r[1], mo))
        docs.append({'cat': str(r[0]).strip(), 'name': str(r[1]).strip(), 'cycle': str(r[2] or '').strip(),
                     'n': int(r[3] or 0), 'marks': marks})
    # 「월별 목록」과 달별 건수 맞춰 보기
    if '월별 목록' in wb.sheetnames:
        lst = Counter()
        for r in list(wb['월별 목록'].iter_rows(values_only=True))[1:]:
            if r and r[0]:
                lst[int(str(r[0]).replace('월', ''))] += 1
        grid = Counter(int(mo) for d in docs for mo in d['marks'])
        if lst != grid:
            raise SystemExit('달력 시트와 월별 목록 건수가 달라요: %s / %s' % (dict(grid), dict(lst)))
    for d in docs:
        if NAME_RE.search(d['name']):
            raise SystemExit('사람 이름처럼 보이는 문서명이 있어요(공개 사이트라 멈춤): 행 %d' % (docs.index(d) + 1))
    cats = [c for c, _ in Counter(d['cat'] for d in docs).most_common()]
    docs.sort(key=lambda d: (cycle_key(d['cycle']), -d['n']))
    return {
        'title': '정례 문서 달력',
        'source': '문서등록대장 2026.3~10월 반복 문서',
        'refreshed_at': m.group(1) if m else '',
        'legend': {'●': '그 달에 기록 있음', '○': '11~2월 예상(매월 문서만, 기록 없음)'},
        'months': MONTHS,
        'cats': cats,
        'cycles': sorted(sorted({d['cycle'] for d in docs}), key=cycle_key),
        'docs': docs,
    }


def main(argv):
    src = argv[1] if len(argv) > 1 else SRC
    out = argv[2] if len(argv) > 2 else OUT
    data = build(src)
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
        f.write('\n')
    print('정례 문서 %d개, 분류 %d개 → %s' % (len(data['docs']), len(data['cats']), os.path.relpath(out, ROOT)))


if __name__ == '__main__':
    main(sys.argv)
